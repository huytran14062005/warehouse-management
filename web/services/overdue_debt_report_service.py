from decimal import Decimal
from io import BytesIO

from flask_mail import Message
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from web import db, mail
from web.dao.sales_invoice_dao import sales_invoice_remaining_amount_expression
from web.models import Customer, Employee, EmployeeRole, InvoiceStatus, SalesInvoice
from web.services.sales_invoice_service import business_today


EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _recipient_emails():
    """Resolve the approved report mailbox, with manager fallback."""
    from flask import current_app

    configured = current_app.config.get("OVERDUE_REPORT_RECIPIENTS", ())
    recipients = [str(email).strip() for email in configured if str(email).strip()]
    if recipients:
        return list(dict.fromkeys(recipients))

    manager_rows = (
        Employee.query
        .filter(
            Employee.role == EmployeeRole.QUAN_LY,
            Employee.is_approved.is_(True),
            Employee.is_locked.is_(False),
        )
        .with_entities(Employee.email)
        .all()
    )
    return list(dict.fromkeys(row.email.strip() for row in manager_rows if row.email))


def get_overdue_customer_rows(as_of_date=None):
    """Aggregate currently overdue, unpaid invoices once per customer."""
    report_date = as_of_date or business_today()
    remaining_amount = sales_invoice_remaining_amount_expression()

    rows = (
        db.session.query(
            Customer.full_name.label("full_name"),
            Customer.phone.label("phone"),
            Customer.address.label("address"),
            func.count(SalesInvoice.id).label("invoice_count"),
            func.sum(remaining_amount).label("overdue_amount"),
            func.min(SalesInvoice.payment_due_date).label("oldest_due_date"),
            func.group_concat(SalesInvoice.id).label("invoice_ids"),
        )
        .join(SalesInvoice, SalesInvoice.customer_id == Customer.id)
        .filter(
            SalesInvoice.status == InvoiceStatus.COMPLETED,
            SalesInvoice.payment_due_date.isnot(None),
            SalesInvoice.payment_due_date < report_date,
            remaining_amount > 0,
        )
        .group_by(Customer.id, Customer.full_name, Customer.phone, Customer.address)
        .order_by(func.sum(remaining_amount).desc(), Customer.full_name.asc())
        .all()
    )

    return [
        {
            "full_name": row.full_name,
            "phone": row.phone,
            "address": row.address,
            "invoice_count": int(row.invoice_count or 0),
            "overdue_amount": Decimal(str(row.overdue_amount or 0)),
            "oldest_due_date": row.oldest_due_date,
            "days_overdue": max(0, (report_date - row.oldest_due_date).days),
            "invoice_ids": str(row.invoice_ids or ""),
        }
        for row in rows
    ]


def _build_workbook(rows, report_date):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Công nợ quá hạn"

    sheet.merge_cells("A1:I1")
    sheet["A1"] = f"DANH SÁCH KHÁCH HÀNG QUÁ HẠN NỢ - {report_date:%d/%m/%Y}"
    sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="123B63")
    sheet["A1"].alignment = Alignment(horizontal="center")

    sheet.merge_cells("A2:I2")
    sheet["A2"] = "Dữ liệu lấy trực tiếp từ các đơn đã hoàn thành còn công nợ tại thời điểm lập báo cáo."
    sheet["A2"].alignment = Alignment(wrap_text=True)

    headers = [
        "STT",
        "Khách hàng",
        "Số điện thoại",
        "Địa chỉ",
        "Số đơn quá hạn",
        "Tổng công nợ quá hạn",
        "Hạn thanh toán sớm nhất",
        "Số ngày quá hạn lâu nhất",
        "Mã đơn quá hạn",
    ]
    header_row = 4
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F75B5")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, row in enumerate(rows, start=1):
        values = [
            index,
            row["full_name"],
            row["phone"],
            row["address"],
            row["invoice_count"],
            float(row["overdue_amount"]),
            row["oldest_due_date"],
            row["days_overdue"],
            row["invoice_ids"],
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(header_row + index, column, value)
        sheet.cell(header_row + index, 6).number_format = '#,##0 "đ"'
        sheet.cell(header_row + index, 7).number_format = "dd/mm/yyyy"

    widths = [8, 28, 18, 42, 16, 22, 22, 24, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:I{max(4, 4 + len(rows))}"
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 34

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def send_overdue_debt_report(as_of_date=None):
    report_date = as_of_date or business_today()
    rows = get_overdue_customer_rows(report_date)
    recipients = _recipient_emails()
    if not recipients:
        raise RuntimeError("Chưa cấu hình email nhận báo cáo công nợ quá hạn")

    message = Message(
        subject=f"Báo cáo khách hàng quá hạn nợ - {report_date:%m/%Y}",
        recipients=recipients,
        body=(
            f"Báo cáo công nợ quá hạn đến ngày {report_date:%d/%m/%Y}.\n\n"
            f"Có {len(rows)} khách hàng còn công nợ quá hạn.\n"
            "Chi tiết nằm trong file Excel đính kèm."
        ),
    )
    message.attach(
        f"bao-cao-cong-no-qua-han-{report_date:%Y-%m}.xlsx",
        EXCEL_CONTENT_TYPE,
        _build_workbook(rows, report_date),
    )
    mail.send(message)
    return {
        "report_date": report_date.isoformat(),
        "customer_count": len(rows),
        "recipients": recipients,
    }
